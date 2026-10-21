import hashlib
from pathlib import Path
from langchain_core.documents import Document
from langchain_text_splitters import RecursiveCharacterTextSplitter
from langchain_community.document_loaders import PyPDFLoader, TextLoader
from docx import Document as DocxDocument


SUPPORTED = {".pdf", ".txt", ".md", ".docx"}

PARENT_SIZE, PARENT_OVERLAP = 2000, 200
CHILD_SIZE, CHILD_OVERLAP = 400, 60


def load_file(path):
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return PyPDFLoader(str(path)).load()
    if suffix in {".txt", ".md"}:
        return TextLoader(str(path), autodetect_encoding=True).load()
    if suffix == ".docx":
        doc = DocxDocument(str(path))
        text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        return [Document(page_content=text, metadata={"source": str(path)})]
    raise ValueError(f"Unsupported file type: {suffix}")


def chunk_documents(docs):
    parent_splitter = RecursiveCharacterTextSplitter(
        chunk_size=PARENT_SIZE, chunk_overlap=PARENT_OVERLAP, add_start_index=True
    )
    child_splitter = RecursiveCharacterTextSplitter(
        chunk_size=CHILD_SIZE, chunk_overlap=CHILD_OVERLAP, add_start_index=True
    )

    parents = {}
    children = []

    for parent_idx, parent in enumerate(parent_splitter.split_documents(list(docs))):
        source = parent.metadata.get("source", "")
        file_name = Path(source).name
        page = parent.metadata.get("page")  # PDFs only, 0-indexed
        parent_start = parent.metadata.get("start_index", 0)
        parent_id = f"{hashlib.md5(source.encode()).hexdigest()[:8]}-p{parent_idx}"

        # --- parent metadata ---
        parent.metadata.update({
            "parent_id": parent_id,
            "file_name": file_name,
            "chunk_level": "parent",
        })
        parents[parent_id] = parent

        # --- metadata injection header (embedded together with the child text) ---
        header_parts = [f"File: {file_name}"]
        if page is not None:
            header_parts.append(f"Page: {page + 1}")
        header = "[" + " | ".join(header_parts) + "]\n"

        # --- children ---
        for child_idx, child in enumerate(child_splitter.split_documents([parent])):
            if not child.page_content.strip():
                continue
            child.metadata.update({
                "child_id": f"{parent_id}-c{child_idx}",
                "child_index": child_idx,
                "chunk_level": "child",
                "start_index": parent_start + child.metadata.get("start_index", 0),
            })
            child.page_content = header + child.page_content
            children.append(child)

    return children, parents


def resolve_parents(child_hits, parents):
    seen, result = set(), []
    for c in child_hits:
        pid = c.metadata["parent_id"]
        if pid not in seen:
            seen.add(pid)
            result.append(parents[pid])
    return result